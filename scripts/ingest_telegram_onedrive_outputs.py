#!/usr/bin/env python3
"""Preserve and audit public OneDrive outputs linked from Telegram @robiai33.

The linked files are derived research outputs, not official judgment originals.
This script preserves them byte-for-byte, creates review-only leads, and never
changes search eligibility, public downloads, platform imports, or a database.
"""
from __future__ import annotations

import argparse
import csv
import hashlib
import json
import re
import shutil
import unicodedata
from collections import Counter, defaultdict
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
TELEGRAM_CHANNEL = "robiai33"
TELEGRAM_POST_ID = 1302
TELEGRAM_POST_URL = "https://t.me/robiai33/1302"
ONEDRIVE_URL = "https://1drv.ms/f/c/987dd35a60362521/IgDYccb6-z8bQbbBPyiFWFIDAccdvrtduwoZA0WOC48ofHo?e=Z84Va1"
FILES = [
    "السوابق_نص_كامل.jsonl",
    "تقرير_المعالجة.csv",
    "فهرس_السوابق_القضائية.xlsx",
    "محرك_السوابق.html",
    "مواصفة_محرك_السوابق.md",
]
RAW_JSONL = "السوابق_نص_كامل.jsonl"
TARGET_DIR = ROOT / "originals" / "telegram-linked" / TELEGRAM_CHANNEL / "post-1302-onedrive"
FILE_REGISTER = ROOT / "manifests" / "collector" / "telegram-linked-file-register.csv"
PENDING = ROOT / "indices" / "collector" / "pending-telegram-linked-precedents.ndjson"
EXACT_DUPLICATES = ROOT / "manifests" / "collector" / "telegram-linked-exact-text-duplicates.csv"
REFERENCE_REVIEW = ROOT / "manifests" / "collector" / "telegram-linked-reference-review.csv"
SUMMARY = ROOT / "manifests" / "collector" / "telegram-linked-import-summary.json"
VALIDATION = ROOT / "manifests" / "collector" / "telegram-linked-validation-report.json"
TELEGRAM_REGISTER = ROOT / "manifests" / "collector" / "telegram-source-register.csv"

DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩۰۱۲۳۴۵۶۷۸۹", "01234567890123456789")
DIACRITICS = re.compile(r"[\u0610-\u061A\u064B-\u065F\u0670\u06D6-\u06ED]")
BIDI = re.compile(r"[\u200E\u200F\u202A-\u202E\u2066-\u2069\uFEFF]")


def normalise(value: Any) -> str:
    text = unicodedata.normalize("NFKC", str(value or "")).translate(DIGITS)
    text = BIDI.sub("", DIACRITICS.sub("", text).replace("ـ", ""))
    return re.sub(r"\s+", " ", text).strip()


def normalise_filename(value: Any) -> str:
    text = normalise(value).casefold()
    return re.sub(r"\s+", " ", text)


def sha256_bytes(value: bytes) -> str:
    return hashlib.sha256(value).hexdigest()


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def sha256_text(value: Any) -> str:
    return sha256_bytes(normalise(value).encode("utf-8"))


def read_ndjson(path: Path) -> list[dict[str, Any]]:
    rows: list[dict[str, Any]] = []
    for number, line in enumerate(path.read_text(encoding="utf-8-sig").splitlines(), start=1):
        if not line.strip():
            continue
        try:
            value = json.loads(line)
        except json.JSONDecodeError as exc:
            raise ValueError(f"Invalid JSON at {path}:{number}: {exc}") from exc
        if not isinstance(value, dict):
            raise ValueError(f"Non-object JSON at {path}:{number}")
        rows.append(value)
    return rows


def write_ndjson(path: Path, rows: list[dict[str, Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        for row in rows:
            handle.write(json.dumps(row, ensure_ascii=False, separators=(",", ":")) + "\n")
    temporary.replace(path)


def write_csv(path: Path, header: list[str], rows: list[list[Any]]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.name + ".tmp")
    with temporary.open("w", encoding="utf-8", newline="") as handle:
        writer = csv.writer(handle, lineterminator="\n")
        writer.writerow(header)
        writer.writerows(rows)
    temporary.replace(path)


def existing_archive_maps() -> tuple[dict[str, set[str]], dict[str, set[str]], set[str]]:
    by_text: dict[str, set[str]] = defaultdict(set)
    by_reference: dict[str, set[str]] = defaultdict(set)
    identifiers: set[str] = set()
    for path in sorted((ROOT / "indices").rglob("*.ndjson")):
        if path == PENDING:
            continue
        for number, line in enumerate(path.read_text(encoding="utf-8", errors="replace").splitlines(), start=1):
            if not line.strip():
                continue
            try:
                row = json.loads(line)
            except json.JSONDecodeError:
                continue
            if not isinstance(row, dict):
                continue
            identifier = str(row.get("id") or row.get("record_id") or f"{path.name}:{number}")
            identifiers.add(identifier)
            for key in ("textChecksum", "text_sha256", "text_checksum"):
                value = normalise(row.get(key))
                if value:
                    by_text[value].add(identifier)
            for key in ("lawsuitNumber", "caseNumber", "case_number", "رقم القضية"):
                value = normalise(row.get(key))
                if value:
                    by_reference[value].add(identifier)
    return by_text, by_reference, identifiers


def telegram_filename_map() -> dict[str, list[dict[str, str]]]:
    matches: dict[str, list[dict[str, str]]] = defaultdict(list)
    with TELEGRAM_REGISTER.open(encoding="utf-8", newline="") as handle:
        for row in csv.DictReader(handle):
            matches[normalise_filename(row["filename_as_published"])].append(row)
    return matches


def xlsx_summary(path: Path) -> dict[str, Any]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    sheets = []
    for sheet in workbook.worksheets:
        sheets.append({"name": sheet.title, "rows": sheet.max_row, "columns": sheet.max_column})
    first_sheet = workbook[workbook.sheetnames[0]]
    header = [value for value in next(first_sheet.iter_rows(min_row=1, max_row=1, values_only=True))]
    workbook.close()
    return {"sheets": sheets, "first_sheet_header": header}


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--input-dir", type=Path, required=True)
    parser.add_argument("--apply", action="store_true")
    parser.add_argument("--output-summary", type=Path)
    args = parser.parse_args()

    missing = [name for name in FILES if not (args.input_dir / name).is_file()]
    if missing:
        raise SystemExit(f"Missing linked files: {missing}")

    acquired_at = datetime.now(timezone(timedelta(hours=3))).replace(microsecond=0).isoformat()
    prior_acquisition: dict[tuple[str, str], str] = {}
    if FILE_REGISTER.exists():
        with FILE_REGISTER.open(encoding="utf-8", newline="") as handle:
            for row in csv.DictReader(handle):
                prior_acquisition[(row.get("filename_as_received", ""), row.get("sha256", ""))] = row.get("acquired_at", "")
    file_rows: list[list[Any]] = []
    file_info: list[dict[str, Any]] = []
    for name in FILES:
        path = args.input_dir / name
        info = {
            "filename": name,
            "bytes": path.stat().st_size,
            "sha256": sha256_file(path),
            "mime_hint": path.suffix.lower().lstrip(".") or "unknown",
        }
        received_at = prior_acquisition.get((name, info["sha256"])) or acquired_at
        file_info.append(info)
        file_rows.append([
            TELEGRAM_CHANNEL,
            TELEGRAM_POST_ID,
            TELEGRAM_POST_URL,
            ONEDRIVE_URL,
            name,
            info["bytes"],
            info["sha256"],
            received_at,
            "public_onedrive_link_from_telegram",
            "derived_research_output",
            "unverified",
            "not_eligible",
            "preserved_for_private_review_not_an_official_legal_original",
        ])

    raw_path = args.input_dir / RAW_JSONL
    source_rows = read_ndjson(raw_path)
    if len(source_rows) != 1213:
        raise ValueError(f"Unexpected JSONL record count: {len(source_rows)}")

    xlsx = xlsx_summary(args.input_dir / "فهرس_السوابق_القضائية.xlsx")
    expected_sheet = next((sheet for sheet in xlsx["sheets"] if sheet["name"] == "فهرس السوابق"), None)
    if expected_sheet is None or expected_sheet["rows"] != len(source_rows) + 1:
        raise ValueError(f"Excel/JSONL record count mismatch: {xlsx['sheets']}")

    telegram_files = telegram_filename_map()
    existing_text, existing_reference, existing_ids = existing_archive_maps()
    seen_ids: set[str] = set(existing_ids)
    first_by_text: dict[str, str] = {}
    pending_rows: list[dict[str, Any]] = []
    exact_rows: list[list[Any]] = []
    reference_rows: list[list[Any]] = []
    source_name_counter: Counter[str] = Counter()
    matched_source_counter: Counter[str] = Counter()
    mode_counter: Counter[str] = Counter()
    nonempty_counter: Counter[str] = Counter()

    for position, source in enumerate(source_rows, start=1):
        full_text = str(source.get("النص الكامل") or "")
        text_hash = sha256_text(full_text)
        source_filename = str(source.get("الملف") or "")
        source_name_counter[source_filename] += 1
        mode = str(source.get("نمط الاستخراج") or "")
        mode_counter[mode] += 1
        matches = telegram_files.get(normalise_filename(source_filename), [])
        if matches:
            matched_source_counter[source_filename] += 1
        source_post_urls = sorted({row["post_url"] for row in matches})
        case_number = normalise(source.get("رقم القضية")) or None
        title = normalise(source.get("الموضوع")) or None
        stable = f"{position}|{source_filename}|{text_hash}"
        identifier = "telegram-linked-precedent-" + sha256_bytes(stable.encode("utf-8"))[:24]
        if identifier in seen_ids:
            raise ValueError(f"Duplicate generated identifier: {identifier}")
        seen_ids.add(identifier)

        for field in ("الموضوع", "رقم القضية", "تاريخ القضية (هـ)", "محكمة الدرجة الأولى", "نتيجة الاستئناف"):
            if normalise(source.get(field)):
                nonempty_counter[field] += 1

        duplicate_ids: set[str] = set(existing_text.get(text_hash, set()))
        if text_hash in first_by_text:
            duplicate_ids.add(first_by_text[text_hash])
        if duplicate_ids:
            exact_rows.append([
                identifier,
                position,
                source_filename,
                text_hash,
                ";".join(sorted(duplicate_ids)),
                "exact_normalized_text",
                "retain_incoming_output_and_review_do_not_promote",
            ])
        else:
            first_by_text[text_hash] = identifier

        related_references = set(existing_reference.get(case_number or "", set()))
        if case_number:
            for previous in pending_rows:
                if previous.get("caseNumber") == case_number:
                    related_references.add(str(previous["id"]))
            if related_references and not duplicate_ids:
                reference_rows.append([
                    identifier,
                    position,
                    source_filename,
                    case_number,
                    text_hash,
                    ";".join(sorted(related_references)),
                    "same_case_reference_distinct_or_unconfirmed_text",
                    "manual_review",
                ])

        duplicate_of = sorted(duplicate_ids)[0] if duplicate_ids else None
        candidate = {
            "id": identifier,
            "documentType": "precedent_candidate",
            "sourceChannel": TELEGRAM_CHANNEL,
            "sourcePostId": TELEGRAM_POST_ID,
            "sourcePostUrl": TELEGRAM_POST_URL,
            "linkedPublicUrl": ONEDRIVE_URL,
            "linkedArtifactFile": f"{TARGET_DIR.relative_to(ROOT).as_posix()}/{RAW_JSONL}",
            "linkedArtifactSha256": next(item["sha256"] for item in file_info if item["filename"] == RAW_JSONL),
            "sourceFilenameAsRecorded": source_filename or None,
            "sourceTelegramPostUrls": source_post_urls,
            "provisionalDocumentClass": "سابقة قضائية مرشحة غير متحققة",
            "sourceRecordNumber": source.get("#") or position,
            "sequenceWithinSource": source.get("م"),
            "title": title,
            "content": full_text,
            "caseNumber": case_number,
            "caseDateHijri": normalise(source.get("تاريخ القضية (هـ)")) or None,
            "firstInstanceCourt": normalise(source.get("محكمة الدرجة الأولى")) or None,
            "appealDecisionNumber": normalise(source.get("رقم قرار الاستئناف")) or None,
            "appealDecisionDateHijri": normalise(source.get("تاريخ القرار (هـ)")) or None,
            "appealCourt": normalise(source.get("محكمة الاستئناف")) or None,
            "appealOutcome": normalise(source.get("نتيجة الاستئناف")) or None,
            "reportedSourcePage": source.get("الصفحة") or None,
            "reportedSourcePageCount": source.get("عدد الصفحات") or None,
            "extractionPattern": mode or None,
            "evidentiaryRankAsReceived": normalise(source.get("مستوى الحجية")) or None,
            "lawCurrencyStatusAsReceived": normalise(source.get("حالة السريان")) or None,
            "textChecksum": text_hash,
            "textStoredInLinkedArtifact": True,
            "officialOriginalBinaryStatus": "missing",
            "officialSourceStatus": "unverified",
            "pageRangeStatus": "unverified_missing_source_pdf",
            "reviewStatus": "incoming_derived_record_pending_source_and_metadata_review",
            "searchIndexEligibility": "not_eligible",
            "publicDownloadEligibility": "not_eligible",
            "platformImportEligibility": "not_eligible",
            "duplicateReviewStatus": "exact_text_candidate" if duplicate_ids else (
                "reference_collision_candidate" if related_references else "no_exact_match_found"
            ),
            "missingRequirements": [
                "official_source_url",
                "official_original_pdf",
                "verified_title" if not title else "",
                "verified_page_range",
                "verified_legal_type",
                "verified_jurisdiction",
            ],
        }
        candidate["missingRequirements"] = [value for value in candidate["missingRequirements"] if value]
        if not duplicate_of:
            pending_rows.append(candidate)

    summary = {
        "schema_version": "1.0",
        "scope": "telegram_public_link_private_intake",
        "telegram_channel": TELEGRAM_CHANNEL,
        "telegram_post_id": TELEGRAM_POST_ID,
        "telegram_post_url": TELEGRAM_POST_URL,
        "linked_public_url": ONEDRIVE_URL,
        "acquired_at": acquired_at,
        "linked_files": file_info,
        "nonempty_linked_files": sum(1 for item in file_info if item["bytes"] > 0),
        "empty_linked_files": [item["filename"] for item in file_info if item["bytes"] == 0],
        "jsonl_records": len(source_rows),
        "unique_text_review_records": len(pending_rows),
        "jsonl_unique_source_filenames": len(source_name_counter),
        "jsonl_source_filenames_matching_telegram_inventory": len(matched_source_counter),
        "jsonl_source_records_matching_telegram_inventory": sum(matched_source_counter.values()),
        "extraction_patterns": dict(sorted(mode_counter.items())),
        "field_completeness": dict(nonempty_counter),
        "xlsx": xlsx,
        "exact_text_duplicate_candidates": len(exact_rows),
        "reference_collision_candidates": len(reference_rows),
        "final_legal_records_added": 0,
        "standalone_legal_pdfs_added": 0,
        "search_eligible_records_added": 0,
        "public_downloads_enabled": False,
        "platform_database_modified": False,
        "classification": "derived_research_outputs_and_provisional_precedent_leads",
        "next_step": "acquire each named source PDF through an authorized Telegram export or official source, then verify SHA-256, type, metadata, and page boundaries",
    }

    validation = {
        "status": "passed",
        "errors": [],
        "files_expected": len(FILES),
        "files_present": len(FILES),
        "files_nonempty": summary["nonempty_linked_files"],
        "jsonl_valid_records": len(source_rows),
        "pending_unique_text_records": len(pending_rows),
        "xlsx_index_rows_excluding_header": expected_sheet["rows"] - 1,
        "jsonl_xlsx_count_match": expected_sheet["rows"] - 1 == len(source_rows),
        "generated_ids_unique": len({row["id"] for row in pending_rows}) == len(pending_rows),
        "all_not_search_eligible": all(row["searchIndexEligibility"] == "not_eligible" for row in pending_rows),
        "all_missing_official_original": all(row["officialOriginalBinaryStatus"] == "missing" for row in pending_rows),
        "final_legal_records_added": 0,
        "platform_database_modified": False,
    }

    if args.apply:
        TARGET_DIR.mkdir(parents=True, exist_ok=True)
        for name in FILES:
            source = args.input_dir / name
            destination = TARGET_DIR / name
            if destination.exists() and sha256_file(destination) != sha256_file(source):
                raise ValueError(f"Refusing to overwrite different preserved file: {destination}")
            if not destination.exists():
                shutil.copyfile(source, destination)
        write_csv(FILE_REGISTER, [
            "telegram_channel", "telegram_post_id", "telegram_post_url", "linked_public_url",
            "filename_as_received", "bytes", "sha256", "acquired_at", "acquisition_method",
            "artifact_class", "official_source_status", "search_eligibility", "notes",
        ], file_rows)
        write_ndjson(PENDING, pending_rows)
        write_csv(EXACT_DUPLICATES, [
            "candidate_id", "source_record_number", "source_filename", "text_sha256",
            "related_record_ids", "match_method", "disposition",
        ], exact_rows)
        write_csv(REFERENCE_REVIEW, [
            "candidate_id", "source_record_number", "source_filename", "case_number",
            "text_sha256", "related_record_ids", "reason", "disposition",
        ], reference_rows)
        SUMMARY.write_text(json.dumps(summary, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
        VALIDATION.write_text(json.dumps(validation, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")

    output = json.dumps({"summary": summary, "validation": validation}, ensure_ascii=False, indent=2) + "\n"
    if args.output_summary:
        args.output_summary.parent.mkdir(parents=True, exist_ok=True)
        args.output_summary.write_text(output, encoding="utf-8")
    print(output, end="")


if __name__ == "__main__":
    main()
